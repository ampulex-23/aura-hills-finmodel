# -*- coding: utf-8 -*-
import sys, re
sys.stdout.reconfigure(encoding='utf-8')
import formulas
from openpyxl import load_workbook
wb = load_workbook('out/aura-hills-tz-simple.xlsx')
gy = wb['Годы']
labels = {}
for row in gy.iter_rows(min_col=1, max_col=1):
    v = row[0].value
    if v: labels[row[0].row] = v
xl = formulas.ExcelModel().loads('out/aura-hills-tz-simple.xlsx').finish()
sol = xl.calculate()
book = 'aura-hills-tz-simple.xlsx'
key_re = re.compile(r"^'\[" + re.escape(book) + r"\]([^']+)'!(\$?[A-Z]+\$?\d+)$")
calc = {}
for k, v in sol.items():
    m = key_re.match(k)
    if not m: continue
    try: val = v.value[0, 0]
    except Exception:
        try: val = v.value
        except Exception: continue
    calc[(m.group(1), m.group(2).replace('$',''))] = val
for rr, lbl in sorted(labels.items()):
    vals = []
    for y in range(5):
        a = f'{chr(66+y)}{rr}'
        v = calc.get(('ГОДЫ', a))
        vals.append(round(v,1) if isinstance(v,(int,float)) else v)
    print(f'{rr:3} {lbl[:52]:54} {vals}')
